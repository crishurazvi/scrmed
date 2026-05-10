from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


@dataclass
class TreatmentRow:
    medicament: str = ""
    dimineata: str = ""
    pranz: str = ""
    seara: str = ""
    observatii: str = ""


@dataclass
class LetterData:
    furnizor: str = "Institutul de Boli Cardiovasculare TIMIȘOARA"
    medic: str = "Dr. OLARIU IOAN"
    contract: str = "II/AS/6/2023"
    casa: str = "CAS Timiș"
    pacient_nume: str = ""
    data_nastere: str = ""
    cnp: str = ""
    data_consult: str = ""
    nr_registru: str = ""
    motiv: str = "Evaluare cardiologică"
    pacient_oncologic: str = "NU"
    diagnostice: List[str] = field(default_factory=list)
    diagnostice_manual: str = ""
    anamneza: str = ""
    factori_risc: str = ""
    ta: str = ""
    fc: str = ""
    examen_local: str = ""
    laborator_normal: str = ""
    laborator_patologic: str = ""
    ekg: str = "EKG anexat"
    eco: str = "ECO"
    rx: str = ""
    alte_investigatii: str = ""
    tratament_efectuat: str = ""
    alte_informatii: str = "Scrisoare medicală valabilă 6 luni"
    tratament: List[TreatmentRow] = field(default_factory=list)
    recomandari: str = "Control peste 6 luni, cu programare."
    prescriptie_status: str = "Nu s-a eliberat prescripție medicală"
    concediu_status: str = "Nu s-a eliberat concediu medical la externare"
    ingrijiri_status: str = "Nu s-a eliberat recomandare pentru îngrijiri medicale la domiciliu/paliative la domiciliu, deoarece nu a fost necesar"
    dispozitive_status: str = "Nu s-a eliberat prescripție medicală pentru dispozitive medicale în ambulatoriu deoarece nu a fost necesar"
    cale_transmitere: str = "prin asigurat"


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def combine_diagnoses(data: LetterData) -> str:
    parts = [d.strip() for d in data.diagnostice if d.strip()]
    if data.diagnostice_manual.strip():
        parts.append(data.diagnostice_manual.strip())
    return ". ".join(parts).strip(" .")


def treatment_rows_from_dataframe(df) -> List[TreatmentRow]:
    rows: List[TreatmentRow] = []
    for _, row in df.iterrows():
        med = normalize_text(row.get("Medicament"))
        if not med:
            continue
        rows.append(
            TreatmentRow(
                medicament=med,
                dimineata=normalize_text(row.get("Dimineața")),
                pranz=normalize_text(row.get("Prânz")),
                seara=normalize_text(row.get("Seara")),
                observatii=normalize_text(row.get("Observații")),
            )
        )
    return rows


def build_plain_letter(data: LetterData) -> str:
    oncologic = "DA" if data.pacient_oncologic == "DA" else "NU"
    diagnostic_text = combine_diagnoses(data) or "[COMPLETARE MANUALĂ: diagnostice confirmate de medic]"
    treatment_lines = []
    if data.tratament:
        treatment_lines.append("Medicament | Dimineața | Prânz | Seara | Observații")
        treatment_lines.append("--- | --- | --- | --- | ---")
        for tr in data.tratament:
            treatment_lines.append(
                f"{tr.medicament} | {tr.dimineata or '-'} | {tr.pranz or '-'} | {tr.seara or '-'} | {tr.observatii}"
            )
    else:
        treatment_lines.append("[COMPLETARE MANUALĂ: tratament recomandat confirmat de medic]")

    return f"""ANEXA 43
Denumire Furnizor: {data.furnizor}
Medic: {data.medic}
Contract/convenție nr.: {data.contract}
{data.casa}

SCRISOARE MEDICALĂ*)

Stimate(ă) coleg(ă), vă informăm că {data.pacient_nume or '[NUME PACIENT]'}, născut la data de {data.data_nastere or '[DATA NAȘTERII]'} CNP/cod unic de Asigurare {data.cnp or '[CNP/COD]'}, a fost consultat în serviciul nostru la data de {data.data_consult or '[DATA CONSULT]'}, / a fost internat în perioada ………… nr. F.O./nr. din Registrul de consultații {data.nr_registru or '[NR.]'}.

Motivele prezentării:
{data.motiv or '[COMPLETARE MANUALĂ]'}

Diagnosticul și codul de diagnostic:
{diagnostic_text}

Pacient oncologic: {oncologic}

Anamneză:
{data.anamneza or '[COMPLETARE MANUALĂ]'}

- factori de risc: {data.factori_risc or '[COMPLETARE MANUALĂ]'}

Examen clinic:
- general: TA={data.ta or '[...]'} mmHg, FC={data.fc or '[...]'} b/min
- local: {data.examen_local or '[COMPLETARE MANUALĂ]'}

Examene de laborator:
- cu valori normale: {data.laborator_normal or '[COMPLETARE MANUALĂ]'}
- cu valori patologice: {data.laborator_patologic or '[COMPLETARE MANUALĂ]'}

Examene paraclinice:
{data.ekg or 'EKG'}
{data.eco or 'ECO'}
Rx: {data.rx or '[COMPLETARE MANUALĂ]'}
Altele: {data.alte_investigatii or '[COMPLETARE MANUALĂ]'}

Tratament efectuat:
{data.tratament_efectuat or '[COMPLETARE MANUALĂ]'}

Alte informații referitoare la starea de sănătate a asiguratului:
{data.alte_informatii or '[COMPLETARE MANUALĂ]'}

Tratament recomandat:
{chr(10).join(treatment_lines)}

Notă: Se va specifica durata pentru care se poate prescrie de medicul din ambulatoriu, inclusiv medicul de familie, fiecare dintre medicamentele recomandate.

Valabilitatea scrisorii medicale începe de la data eliberării acesteia.
Valabilitatea este în concordanță cu protocolul terapeutic.
În cazul în care medicul de specialitate nu consemnează o valabilitate pentru conduita terapeutică recomandată, valabilitatea scrisorii medicale încetează în momentul în care medicul de familie recomandă pacientului reevaluarea stării de sănătate.

Indicație de revenire pentru internare:
☐ da, revine pentru internare în termen de ........................................................
☒ nu, nu este necesară revenirea pentru internare

{data.prescriptie_status}
{data.concediu_status}
{data.ingrijiri_status}
{data.dispozitive_status}

Data: {data.data_consult or str(date.today())}                                      Semnătura și parafa medicului
                                                                      .............................

{data.recomandari or 'Control peste 6 luni, cu programare.'}

Calea de transmitere: {data.cale_transmitere}
"""


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)


def add_label_value(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(value or "[COMPLETARE MANUALĂ]")


def build_docx(data: LetterData) -> bytes:
    doc = Document()
    set_doc_defaults(doc)

    p = doc.add_paragraph("ANEXA 43")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_paragraph(f"Denumire Furnizor: {data.furnizor}")
    doc.add_paragraph(f"Medic: {data.medic}")
    doc.add_paragraph(f"Contract/convenție nr.: {data.contract}")
    doc.add_paragraph(data.casa)

    title = doc.add_paragraph("SCRISOARE MEDICALĂ*)")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True

    intro = (
        f"Stimate(ă) coleg(ă), vă informăm că {data.pacient_nume or '[NUME PACIENT]'}, "
        f"născut la data de {data.data_nastere or '[DATA NAȘTERII]'} "
        f"CNP/cod unic de Asigurare {data.cnp or '[CNP/COD]'}, "
        f"a fost consultat în serviciul nostru la data de {data.data_consult or '[DATA CONSULT]'}, "
        f"/ a fost internat în perioada ………… nr. F.O./nr. din Registrul de consultații {data.nr_registru or '[NR.]'}."
    )
    doc.add_paragraph(intro)

    add_label_value(doc, "Motivele prezentării:\n", data.motiv)
    add_label_value(doc, "Diagnosticul și codul de diagnostic:\n", combine_diagnoses(data))
    doc.add_paragraph(f"Pacient oncologic: {'☒ DA   ☐ NU' if data.pacient_oncologic == 'DA' else '☐ DA   ☒ NU'}")
    add_label_value(doc, "Anamneză:\n", data.anamneza)
    add_label_value(doc, "- factori de risc: ", data.factori_risc)

    doc.add_paragraph("Examen clinic:").runs[0].bold = True
    doc.add_paragraph(f"- general: TA={data.ta or '[...]'} mmHg, FC={data.fc or '[...]'} b/min")
    doc.add_paragraph(f"- local: {data.examen_local or '[COMPLETARE MANUALĂ]'}")

    doc.add_paragraph("Examene de laborator:").runs[0].bold = True
    doc.add_paragraph(f"- cu valori normale: {data.laborator_normal or '[COMPLETARE MANUALĂ]'}")
    doc.add_paragraph(f"- cu valori patologice: {data.laborator_patologic or '[COMPLETARE MANUALĂ]'}")

    doc.add_paragraph("Examene paraclinice:").runs[0].bold = True
    doc.add_paragraph(data.ekg or "EKG")
    doc.add_paragraph(data.eco or "ECO")
    doc.add_paragraph(f"Rx: {data.rx or '[COMPLETARE MANUALĂ]'}")
    doc.add_paragraph(f"Altele: {data.alte_investigatii or '[COMPLETARE MANUALĂ]'}")

    add_label_value(doc, "Tratament efectuat:\n", data.tratament_efectuat)
    add_label_value(doc, "Alte informații referitoare la starea de sănătate a asiguratului:\n", data.alte_informatii)

    doc.add_paragraph("Tratament recomandat:").runs[0].bold = True
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Medicament", "Dimineața", "Prânz", "Seara", "Observații"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    if data.tratament:
        for tr in data.tratament:
            row = table.add_row().cells
            row[0].text = tr.medicament
            row[1].text = tr.dimineata or "-"
            row[2].text = tr.pranz or "-"
            row[3].text = tr.seara or "-"
            row[4].text = tr.observatii
    else:
        row = table.add_row().cells
        row[0].text = "[COMPLETARE MANUALĂ]"
        row[1].text = "-"
        row[2].text = "-"
        row[3].text = "-"
        row[4].text = ""

    doc.add_paragraph(
        "Notă: Se va specifica durata pentru care se poate prescrie de medicul din ambulatoriu, inclusiv medicul de familie, fiecare dintre medicamentele recomandate."
    )
    doc.add_paragraph("Valabilitatea scrisorii medicale începe de la data eliberării acesteia.")
    doc.add_paragraph("Valabilitatea este în concordanță cu protocolul terapeutic.")
    doc.add_paragraph(
        "În cazul în care medicul de specialitate nu consemnează o valabilitate pentru conduita terapeutică recomandată, valabilitatea scrisorii medicale încetează în momentul în care medicul de familie recomandă pacientului reevaluarea stării de sănătate."
    )

    doc.add_paragraph("Indicație de revenire pentru internare")
    doc.add_paragraph("☐ da, revine pentru internare în termen de ........................................................")
    doc.add_paragraph("☒ nu, nu este necesară revenirea pentru internare")
    doc.add_paragraph(data.prescriptie_status)
    doc.add_paragraph(data.concediu_status)
    doc.add_paragraph(data.ingrijiri_status)
    doc.add_paragraph(data.dispozitive_status)

    doc.add_paragraph(f"Data: {data.data_consult or str(date.today())}                                      Semnătura și parafa medicului")
    doc.add_paragraph("                                                                      .............................")
    doc.add_paragraph(data.recomandari or "Control peste 6 luni, cu programare.")
    doc.add_paragraph(f"Calea de transmitere: {data.cale_transmitere}")

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def build_pdf(data: LetterData) -> bytes:
    """Simple optional PDF export. DOCX remains the preferred output."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas

    text = build_plain_letter(data).replace("☒", "[x]").replace("☐", "[ ]")
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 1.5 * cm
    x = 1.5 * cm
    c.setFont("Times-Roman", 9)
    for raw_line in text.splitlines():
        line = raw_line
        while len(line) > 115:
            c.drawString(x, y, line[:115])
            line = line[115:]
            y -= 0.45 * cm
            if y < 1.5 * cm:
                c.showPage(); c.setFont("Times-Roman", 9); y = height - 1.5 * cm
        c.drawString(x, y, line)
        y -= 0.45 * cm
        if y < 1.5 * cm:
            c.showPage(); c.setFont("Times-Roman", 9); y = height - 1.5 * cm
    c.save()
    return buffer.getvalue()
