import json
import html
from io import BytesIO
from datetime import date, datetime
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components

try:
    from docx import Document
    DOCX_AVAILABLE = True
except Exception:
    Document = None
    DOCX_AVAILABLE = False


APP_TITLE = "Anexa 43 – Scrisoare medicală"


def init_state() -> None:
    defaults = {
        "diagnostice": [],
        "tratamente": [],
        "investigatii_selectate": [],
        "scrisoare_generata": "",
        "last_json_path": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def format_date(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if value:
        return str(value)
    return ""


def v(value, placeholder: str = "necompletat") -> str:
    """Returnează text curat sau un placeholder discret pentru câmpuri goale."""
    if value is None:
        return placeholder
    text = str(value).strip()
    return text if text else placeholder


def join_list(items, placeholder: str = "necompletat") -> str:
    items = [str(x).strip() for x in items if str(x).strip()]
    return ", ".join(items) if items else placeholder


def diagnosis_label(diag: dict) -> str:
    text = diag.get("text", "").strip()
    cim = diag.get("cim10", "").strip()
    return f"{text} [{cim}]" if cim else text


def build_treatment_text(treatments: list[dict]) -> str:
    if not treatments:
        return "necompletat"

    lines = ["Medicament | Dimineața | Prânz | Seara | Durată/observații"]
    lines.append("-" * 72)
    for row in treatments:
        med = v(row.get("medicament", ""), "")
        dim = v(row.get("dimineata", ""), "-")
        pranz = v(row.get("pranz", ""), "-")
        seara = v(row.get("seara", ""), "-")
        obs = v(row.get("observatii", ""), "-")
        lines.append(f"{med} | {dim} | {pranz} | {seara} | {obs}")
    return "\n".join(lines)


def build_checkbox_line(label: str, selected: bool) -> str:
    return f"{'☒' if selected else '☐'} {label}"


def build_letter(data: dict) -> str:
    diagnostice = st.session_state.diagnostice
    tratamente = st.session_state.tratamente

    if diagnostice:
        diagnostic_principal = diagnosis_label(diagnostice[0])
        diagnostice_secundare = diagnostice[1:]
    else:
        diagnostic_principal = "necompletat"
        diagnostice_secundare = []

    if diagnostice_secundare:
        diag_sec_text = "\n".join(
            f"{idx}. {diagnosis_label(diag)}"
            for idx, diag in enumerate(diagnostice_secundare, start=1)
        )
    else:
        diag_sec_text = "necompletat"

    antecedente_items = []
    antecedente_items.extend(data.get("antecedente_selectate", []))
    if data.get("antecedente_text"):
        antecedente_items.append(data["antecedente_text"])

    factori_risc_items = []
    factori_risc_items.extend(data.get("factori_risc_selectati", []))
    if data.get("factori_risc_text"):
        factori_risc_items.append(data["factori_risc_text"])

    investigatii_items = []
    investigatii_items.extend(st.session_state.investigatii_selectate)
    if data.get("lab_normale"):
        investigatii_items.append(f"Laborator cu valori normale: {data['lab_normale']}")
    if data.get("lab_patologice"):
        investigatii_items.append(f"Laborator cu valori patologice: {data['lab_patologice']}")
    if data.get("ekg"):
        investigatii_items.append(f"EKG: {data['ekg']}")
    if data.get("eco"):
        investigatii_items.append(f"ECO: {data['eco']}")
    if data.get("rx"):
        investigatii_items.append(f"Rx: {data['rx']}")
    if data.get("alte_investigatii"):
        investigatii_items.append(f"Altele: {data['alte_investigatii']}")

    recomandari_items = []
    recomandari_items.extend(data.get("recomandari_selectate", []))
    if data.get("recomandari_text"):
        recomandari_items.append(data["recomandari_text"])

    perioada = ""
    tip_prezentare = data.get("tip_prezentare", "Consultație")
    if tip_prezentare == "Internare":
        perioada = (
            f"a fost internat(ă) în perioada "
            f"{format_date(data.get('data_internare'))} – {format_date(data.get('data_externare'))}"
        )
    else:
        perioada = f"a fost consultat(ă) în serviciul nostru la data de {format_date(data.get('data_consult'))}"

    prescriptie_status = data.get("prescriptie_status", "Nu s-a eliberat prescripție medicală")
    concediu_status = data.get("concediu_status", "Nu s-a eliberat concediu medical")
    ingrijiri_status = data.get("ingrijiri_status", "Nu s-a eliberat recomandare pentru îngrijiri la domiciliu")
    dispozitive_status = data.get("dispozitive_status", "Nu s-a eliberat prescripție pentru dispozitive medicale")

    lines = [
        "ANEXA 43",
        "",
        "SCRISOARE MEDICALĂ",
        "",
        f"Denumire furnizor: {v(data.get('furnizor'))}",
        f"Medic: {v(data.get('medic'))}",
        f"Contract/convenție nr.: {v(data.get('contract'))}",
        f"CAS: {v(data.get('cas'))}",
        "",
        (
            f"Stimate(ă) coleg(ă), vă informăm că {v(data.get('nume_pacient'))}, "
            f"născut(ă) la data de {format_date(data.get('data_nasterii')) or 'necompletat'}, "
            f"CNP/cod unic de asigurare {v(data.get('cnp'))}, {perioada}, "
            f"nr. F.O./nr. din Registrul de consultații: {v(data.get('fo_registru'))}."
        ),
        "",
        "DATE PACIENT",
        f"Nume și prenume: {v(data.get('nume_pacient'))}",
        f"Data nașterii: {format_date(data.get('data_nasterii')) or 'necompletat'}",
        f"Sex: {v(data.get('sex'))}",
        f"CNP/cod unic de asigurare: {v(data.get('cnp'))}",
        f"Adresă: {v(data.get('adresa'))}",
        "",
        "PERIOADA CONSULTAȚIEI / INTERNĂRII",
        f"Tip prezentare: {v(tip_prezentare)}",
        f"Data consultului: {format_date(data.get('data_consult')) or 'necompletat'}",
        f"Data internării: {format_date(data.get('data_internare')) or 'necompletat'}",
        f"Data externării: {format_date(data.get('data_externare')) or 'necompletat'}",
        "",
        "MOTIVELE PREZENTĂRII",
        v(data.get("motiv_prezentare")),
        "",
        "DIAGNOSTIC PRINCIPAL",
        diagnostic_principal,
        "",
        "DIAGNOSTICE SECUNDARE",
        diag_sec_text,
        "",
        "PACIENT ONCOLOGIC",
        v(data.get("pacient_oncologic")),
        "",
        "ANAMNEZĂ",
        v(data.get("anamneza")),
        "",
        "ANTECEDENTE PERSONALE",
        join_list(antecedente_items),
        "",
        "ANTECEDENTE HEREDOCOLATERALE",
        v(data.get("antecedente_heredo")),
        "",
        "FACTORI DE RISC",
        join_list(factori_risc_items),
        "",
        "EXAMEN CLINIC",
        f"General: {v(data.get('examen_general'))}",
        f"Local: {v(data.get('examen_local'))}",
        "",
        "INVESTIGAȚII EFECTUATE",
        join_list(investigatii_items),
        "",
        "TRATAMENT ANTERIOR",
        v(data.get("tratament_anterior")),
        "",
        "TRATAMENT EFECTUAT / ADMINISTRAT",
        v(data.get("tratament_efectuat")),
        "",
        "TRATAMENT RECOMANDAT LA EXTERNARE / DUPĂ CONSULT",
        build_treatment_text(tratamente),
        "",
        "ALTE INFORMAȚII REFERITOARE LA STAREA DE SĂNĂTATE A ASIGURATULUI",
        v(data.get("alte_informatii")),
        "",
        "EVOLUȚIE",
        v(data.get("evolutie")),
        "",
        "CONCLUZII",
        v(data.get("concluzii")),
        "",
        "RECOMANDĂRI",
        join_list(recomandari_items),
        "",
        "INFORMAȚII ADMINISTRATIVE",
        build_checkbox_line("S-a eliberat prescripție medicală", prescriptie_status == "S-a eliberat prescripție medicală"),
        build_checkbox_line("Nu s-a eliberat prescripție medicală deoarece nu a fost necesar", prescriptie_status == "Nu a fost necesară prescripție medicală"),
        build_checkbox_line("Nu s-a eliberat prescripție medicală", prescriptie_status == "Nu s-a eliberat prescripție medicală"),
        f"Seria și numărul prescripției: {v(data.get('serie_prescriptie'))}",
        "",
        build_checkbox_line("S-a eliberat concediu medical", concediu_status == "S-a eliberat concediu medical"),
        build_checkbox_line("Nu s-a eliberat concediu medical deoarece nu a fost necesar", concediu_status == "Nu a fost necesar concediu medical"),
        build_checkbox_line("Nu s-a eliberat concediu medical", concediu_status == "Nu s-a eliberat concediu medical"),
        f"Seria și numărul concediului medical: {v(data.get('serie_concediu'))}",
        "",
        build_checkbox_line("S-a eliberat recomandare pentru îngrijiri medicale/paliative la domiciliu", ingrijiri_status == "S-a eliberat recomandare pentru îngrijiri la domiciliu"),
        build_checkbox_line("Nu s-a eliberat recomandare pentru îngrijiri medicale/paliative la domiciliu deoarece nu a fost necesar", ingrijiri_status == "Nu a fost necesară recomandare pentru îngrijiri la domiciliu"),
        "",
        build_checkbox_line("S-a eliberat prescripție pentru dispozitive medicale în ambulatoriu", dispozitive_status == "S-a eliberat prescripție pentru dispozitive medicale"),
        build_checkbox_line("Nu s-a eliberat prescripție pentru dispozitive medicale în ambulatoriu deoarece nu a fost necesar", dispozitive_status == "Nu a fost necesară prescripție pentru dispozitive medicale"),
        "",
        f"Valabilitate scrisoare medicală: {v(data.get('valabilitate'))}",
        f"Control recomandat: {v(data.get('control'))}",
        f"Calea de transmitere: {v(data.get('cale_transmitere'))}",
        "",
        f"Data: {format_date(data.get('data_emitere')) or 'necompletat'}",
        "",
        "Semnătura și parafa medicului",
        v(data.get("semnatura")),
    ]

    return "\n".join(lines)


def build_payload(data: dict) -> dict:
    safe_data = {}
    for key, value in data.items():
        if isinstance(value, date):
            safe_data[key] = value.isoformat()
        else:
            safe_data[key] = value

    return {
        "metadata": {
            "app": APP_TITLE,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        },
        "fields": safe_data,
        "diagnostice": st.session_state.diagnostice,
        "tratamente": st.session_state.tratamente,
        "investigatii_selectate": st.session_state.investigatii_selectate,
        "scrisoare_generata": st.session_state.scrisoare_generata,
    }


def make_docx_bytes(text: str) -> BytesIO:
    document = Document()
    document.add_heading("ANEXA 43 – SCRISOARE MEDICALĂ", level=1)

    for line in text.splitlines():
        clean = line.strip()
        if not clean:
            document.add_paragraph("")
        elif clean.isupper() and len(clean) < 80:
            document.add_heading(clean, level=2)
        else:
            document.add_paragraph(clean)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def render_copy_component(text: str) -> None:
    js_text = json.dumps(text)
    escaped_text = html.escape(text)
    components.html(
        f"""
        <div style="font-family: sans-serif;">
            <button
                onclick='copyLetterText()'
                style="
                    padding: 0.55rem 0.9rem;
                    border-radius: 0.45rem;
                    border: 1px solid #ccc;
                    background: #f7f7f7;
                    cursor: pointer;
                    margin-bottom: 0.5rem;
                "
            >
                Copiază textul în clipboard
            </button>
            <span id="copy-status" style="margin-left: 0.75rem; color: #2e7d32;"></span>
            <textarea id="letter-copy-box" style="position:absolute; left:-9999px;">{escaped_text}</textarea>
        </div>
        <script>
            function copyLetterText() {{
                const text = {js_text};
                const status = document.getElementById("copy-status");

                if (navigator.clipboard && window.isSecureContext) {{
                    navigator.clipboard.writeText(text).then(function() {{
                        status.innerText = "Copiat.";
                    }}).catch(function() {{
                        fallbackCopy(text, status);
                    }});
                }} else {{
                    fallbackCopy(text, status);
                }}
            }}

            function fallbackCopy(text, status) {{
                const area = document.getElementById("letter-copy-box");
                area.value = text;
                area.focus();
                area.select();
                try {{
                    document.execCommand("copy");
                    status.innerText = "Copiat.";
                }} catch (err) {{
                    status.innerText = "Selectează textul din caseta de mai jos și copiază manual.";
                }}
            }}
        </script>
        """,
        height=70,
    )


def main() -> None:
    st.set_page_config(
        page_title=APP_TITLE,
        page_icon="🩺",
        layout="wide",
    )

    init_state()

    st.title(APP_TITLE)
    st.caption(
        "Generator minimalist pentru completarea rapidă a unei scrisori medicale tip Anexa 43. "
        "Aplicația nu inventează date medicale: câmpurile necompletate rămân marcate discret."
    )

    with st.sidebar:
        st.header("Acțiuni rapide")

        if st.button("Șterge diagnosticele"):
            st.session_state.diagnostice = []
            st.success("Lista de diagnostice a fost ștearsă.")

        if st.button("Șterge tratamentele"):
            st.session_state.tratamente = []
            st.success("Lista de tratamente a fost ștearsă.")

        if st.button("Resetează scrisoarea generată"):
            st.session_state.scrisoare_generata = ""
            st.success("Textul generat a fost resetat.")

        st.divider()
        st.write("Fișiere utile:")
        st.write("- `app.py`")
        st.write("- `requirements.txt`")
        st.info("Pe Streamlit Community Cloud, salvarea locală JSON nu este persistentă pe termen lung.")

    data = {}

    with st.expander("1. Date furnizor / medic", expanded=True):
        c1, c2 = st.columns(2)
        with c1:
            data["furnizor"] = st.text_input("Denumire furnizor")
            data["contract"] = st.text_input("Contract/convenție nr.")
        with c2:
            data["medic"] = st.text_input("Medic")
            data["cas"] = st.text_input("CAS")

    with st.expander("2. Date pacient", expanded=True):
        c1, c2, c3 = st.columns([2, 1, 1])
        with c1:
            data["nume_pacient"] = st.text_input("Nume și prenume pacient")
            data["adresa"] = st.text_input("Adresă")
        with c2:
            data["data_nasterii"] = st.date_input("Data nașterii", value=None, format="DD.MM.YYYY")
            data["sex"] = st.selectbox("Sex", ["", "F", "M", "Altul / nespecificat"])
        with c3:
            data["cnp"] = st.text_input("CNP / cod unic de asigurare")
            data["pacient_oncologic"] = st.selectbox("Pacient oncologic", ["", "Nu", "Da"])

    with st.expander("3. Consultație / internare", expanded=True):
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            data["tip_prezentare"] = st.selectbox("Tip prezentare", ["Consultație", "Internare"])
        with c2:
            data["data_consult"] = st.date_input("Data consultului", value=date.today(), format="DD.MM.YYYY")
        with c3:
            data["data_internare"] = st.date_input("Data internării", value=None, format="DD.MM.YYYY")
        with c4:
            data["data_externare"] = st.date_input("Data externării", value=None, format="DD.MM.YYYY")

        data["fo_registru"] = st.text_input("Nr. F.O. / nr. Registru consultații")
        data["motiv_prezentare"] = st.text_area(
            "Motivele prezentării",
            height=90,
            placeholder="Ex: evaluare cardiologică, dispnee, durere toracică etc.",
        )

    with st.expander("4. Diagnostice", expanded=True):
        st.write("Primul diagnostic introdus va fi considerat diagnostic principal. Următoarele vor fi diagnostice secundare.")

        with st.form("diagnostic_form", clear_on_submit=True):
            c1, c2 = st.columns([3, 1])
            with c1:
                new_diag = st.text_input("Diagnostic nou")
            with c2:
                new_cim = st.text_input("Cod CIM-10, opțional")
            submitted_diag = st.form_submit_button("Adaugă diagnostic")

        if submitted_diag:
            if new_diag.strip():
                st.session_state.diagnostice.append(
                    {
                        "text": new_diag.strip(),
                        "cim10": new_cim.strip(),
                    }
                )
                st.success("Diagnosticul a fost adăugat.")
            else:
                st.warning("Completează diagnosticul înainte de adăugare.")

        if st.session_state.diagnostice:
            st.subheader("Lista diagnosticelor")
            for idx, diag in enumerate(st.session_state.diagnostice):
                role = "Diagnostic principal" if idx == 0 else f"Diagnostic secundar {idx}"
                c1, c2 = st.columns([5, 1])
                with c1:
                    st.write(f"**{role}:** {diagnosis_label(diag)}")
                with c2:
                    if st.button("Șterge", key=f"delete_diag_{idx}"):
                        st.session_state.diagnostice.pop(idx)
                        st.rerun()
        else:
            st.info("Nu există diagnostice introduse încă.")

    with st.expander("5. Anamneză, antecedente și factori de risc"):
        data["anamneza"] = st.text_area("Anamneză", height=100)

        antecedente_options = [
            "Hipertensiune arterială",
            "Diabet zaharat",
            "Dislipidemie",
            "Boală coronariană",
            "Infarct miocardic vechi",
            "Fibrilație atrială",
            "Insuficiență cardiacă",
            "Boală cronică de rinichi",
            "AVC ischemic vechi",
            "BPOC / bronșită cronică",
            "Obezitate",
        ]
        data["antecedente_selectate"] = st.multiselect(
            "Antecedente personale frecvente",
            antecedente_options,
        )
        data["antecedente_text"] = st.text_area("Alte antecedente personale", height=80)
        data["antecedente_heredo"] = st.text_area("Antecedente heredocolaterale", height=80)

        factori_options = [
            "Tabagism activ",
            "Tabagism vechi",
            "HTA",
            "Diabet zaharat",
            "Dislipidemie",
            "Obezitate",
            "Sedentarism",
            "Antecedente familiale cardiovasculare",
        ]
        data["factori_risc_selectati"] = st.multiselect("Factori de risc", factori_options)
        data["factori_risc_text"] = st.text_area("Alți factori de risc", height=80)

    with st.expander("6. Examen clinic"):
        c1, c2 = st.columns(2)
        with c1:
            ta = st.text_input("TA", placeholder="Ex: 140/80 mmHg")
        with c2:
            fc = st.text_input("FC", placeholder="Ex: 70 b/min")

        default_general = []
        if ta.strip():
            default_general.append(f"TA={ta.strip()}")
        if fc.strip():
            default_general.append(f"FC={fc.strip()}")

        data["examen_general"] = st.text_area(
            "Examen clinic general",
            value=", ".join(default_general),
            height=90,
        )
        data["examen_local"] = st.text_area("Examen local", height=80)

    with st.expander("7. Investigații"):
        investigatii_options = [
            "EKG anexat",
            "ECO anexat",
            "Rx toracic",
            "Holter ECG",
            "Holter TA",
            "Test de efort",
            "Coronarografie",
            "Angio-CT",
            "Analize biologice",
        ]
        st.session_state.investigatii_selectate = st.multiselect(
            "Investigații efectuate / anexate",
            investigatii_options,
            default=st.session_state.investigatii_selectate,
            key="investigatii_multiselect",
        )

        c1, c2 = st.columns(2)
        with c1:
            data["lab_normale"] = st.text_area("Examene de laborator cu valori normale", height=80)
            data["ekg"] = st.text_area("EKG", height=80)
            data["rx"] = st.text_area("Rx", height=80)
        with c2:
            data["lab_patologice"] = st.text_area("Examene de laborator cu valori patologice", height=80)
            data["eco"] = st.text_area("ECO", height=80)
            data["alte_investigatii"] = st.text_area("Alte investigații / de efectuat", height=80)

    with st.expander("8. Tratament anterior, efectuat și recomandat", expanded=True):
        data["tratament_anterior"] = st.text_area("Tratament anterior", height=80)
        data["tratament_efectuat"] = st.text_area("Tratament efectuat / administrat", height=80)

        st.subheader("Tratament recomandat")
        st.write("Adaugă medicamentele unul câte unul. Lista este păstrată în `st.session_state`.")

        with st.form("treatment_form", clear_on_submit=True):
            c1, c2, c3, c4, c5 = st.columns([2, 1, 1, 1, 2])
            with c1:
                med = st.text_input("Medicament")
            with c2:
                dim = st.text_input("Dimineața")
            with c3:
                pranz = st.text_input("Prânz")
            with c4:
                seara = st.text_input("Seara")
            with c5:
                obs = st.text_input("Durată / observații")
            submitted_treatment = st.form_submit_button("Adaugă tratament")

        if submitted_treatment:
            if med.strip():
                st.session_state.tratamente.append(
                    {
                        "medicament": med.strip(),
                        "dimineata": dim.strip(),
                        "pranz": pranz.strip(),
                        "seara": seara.strip(),
                        "observatii": obs.strip(),
                    }
                )
                st.success("Tratamentul a fost adăugat.")
            else:
                st.warning("Completează numele medicamentului înainte de adăugare.")

        if st.session_state.tratamente:
            st.write("**Listă tratamente recomandate:**")
            for idx, row in enumerate(st.session_state.tratamente):
                c1, c2 = st.columns([6, 1])
                with c1:
                    st.write(
                        f"{idx + 1}. **{row['medicament']}** | "
                        f"Dimineața: {row.get('dimineata') or '-'} | "
                        f"Prânz: {row.get('pranz') or '-'} | "
                        f"Seara: {row.get('seara') or '-'} | "
                        f"Obs: {row.get('observatii') or '-'}"
                    )
                with c2:
                    if st.button("Șterge", key=f"delete_treatment_{idx}"):
                        st.session_state.tratamente.pop(idx)
                        st.rerun()
        else:
            st.info("Nu există tratamente recomandate introduse încă.")

    with st.expander("9. Evoluție, concluzii și recomandări"):
        data["alte_informatii"] = st.text_area("Alte informații referitoare la starea de sănătate", height=80)
        data["evolutie"] = st.text_area("Evoluție", height=90)
        data["concluzii"] = st.text_area("Concluzii", height=90)

        recomandari_options = [
            "Control cardiologic peste 6 luni",
            "Monitorizare TA la domiciliu",
            "Regim hiposodat",
            "Continuarea tratamentului conform recomandărilor",
            "Analize biologice de control",
            "Prezentare la UPU în caz de agravare",
            "Consult diabetologic",
            "Consult neurologic",
            "Consult gastroenterologic",
            "Consult chirurgie cardiovasculară",
        ]
        data["recomandari_selectate"] = st.multiselect("Recomandări uzuale", recomandari_options)
        data["recomandari_text"] = st.text_area("Alte recomandări", height=100)

    with st.expander("10. Informații administrative"):
        c1, c2 = st.columns(2)
        with c1:
            data["prescriptie_status"] = st.radio(
                "Prescripție medicală",
                [
                    "S-a eliberat prescripție medicală",
                    "Nu a fost necesară prescripție medicală",
                    "Nu s-a eliberat prescripție medicală",
                ],
                index=2,
            )
            data["serie_prescriptie"] = st.text_input("Serie și număr prescripție")

            data["ingrijiri_status"] = st.radio(
                "Îngrijiri medicale/paliative la domiciliu",
                [
                    "S-a eliberat recomandare pentru îngrijiri la domiciliu",
                    "Nu a fost necesară recomandare pentru îngrijiri la domiciliu",
                ],
                index=1,
            )
        with c2:
            data["concediu_status"] = st.radio(
                "Concediu medical",
                [
                    "S-a eliberat concediu medical",
                    "Nu a fost necesar concediu medical",
                    "Nu s-a eliberat concediu medical",
                ],
                index=2,
            )
            data["serie_concediu"] = st.text_input("Serie și număr concediu medical")

            data["dispozitive_status"] = st.radio(
                "Dispozitive medicale în ambulatoriu",
                [
                    "S-a eliberat prescripție pentru dispozitive medicale",
                    "Nu a fost necesară prescripție pentru dispozitive medicale",
                ],
                index=1,
            )

        c3, c4, c5 = st.columns(3)
        with c3:
            data["valabilitate"] = st.text_input("Valabilitate", placeholder="Ex: 6 luni")
        with c4:
            data["control"] = st.text_input("Control recomandat", placeholder="Ex: peste 6 luni")
        with c5:
            data["cale_transmitere"] = st.selectbox("Calea de transmitere", ["prin asigurat", "prin poștă", "poștă electronică", "altă cale"])

        data["data_emitere"] = st.date_input("Data emiterii", value=date.today(), format="DD.MM.YYYY")
        data["semnatura"] = st.text_input("Semnătură / parafă medic")

    st.divider()

    c1, c2, c3 = st.columns([1, 1, 2])
    with c1:
        generate_clicked = st.button("Generează scrisoarea medicală", type="primary")
    with c2:
        save_json_clicked = st.button("Salvează JSON local pe server")

    if generate_clicked:
        st.session_state.scrisoare_generata = build_letter(data)
        st.success("Scrisoarea medicală a fost generată.")

    payload = build_payload(data)
    json_text = json.dumps(payload, ensure_ascii=False, indent=2)

    if save_json_clicked:
        Path("data").mkdir(exist_ok=True)
        filename = f"data/anexa43_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        Path(filename).write_text(json_text, encoding="utf-8")
        st.session_state.last_json_path = filename
        st.success(f"Datele au fost salvate local pe server: {filename}")

    if st.session_state.last_json_path:
        st.caption(f"Ultimul JSON salvat local: `{st.session_state.last_json_path}`")

    st.header("Scrisoare generată")

    if st.session_state.scrisoare_generata:
        st.text_area(
            "Text final, ușor de copiat în Word",
            value=st.session_state.scrisoare_generata,
            height=520,
        )

        render_copy_component(st.session_state.scrisoare_generata)

        st.download_button(
            "Descarcă .txt",
            data=st.session_state.scrisoare_generata.encode("utf-8"),
            file_name="anexa43_scrisoare_medicala.txt",
            mime="text/plain",
        )

        st.download_button(
            "Descarcă datele .json",
            data=json_text.encode("utf-8"),
            file_name="anexa43_date.json",
            mime="application/json",
        )

        if DOCX_AVAILABLE:
            st.download_button(
                "Descarcă .docx",
                data=make_docx_bytes(st.session_state.scrisoare_generata),
                file_name="anexa43_scrisoare_medicala.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        else:
            st.warning("Exportul .docx nu este disponibil. Instalează `python-docx`.")
    else:
        st.info("Completează formularul și apasă „Generează scrisoarea medicală”.")


if __name__ == "__main__":
    main()
